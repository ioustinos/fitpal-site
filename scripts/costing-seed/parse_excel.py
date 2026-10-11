#!/usr/bin/env python3
"""WEC-860 (seed costing data from the kitchen Excel) — step 1: parse.

Reads «Appendix_B_Updated.xlsx» and writes one JSON file with the rows the
seed needs. Nothing here touches a database.

    python3 -I scripts/costing-seed/parse_excel.py <xlsx> <out.json>

Sheets read (values only — formulas are taken at their cached result):
  Src Costs          A category · B raw item name · C price per kg (net)
  Src Packaging      A name · B pieces per pack · C price per pack · D price per piece
  New Index Υλικών   A raw item (→ Src Costs B) · B component name ·
                     G cost/100 g raw · H cooked g per 100 g raw · I cost/100 g cooked
  Group Υλικών       B composite name · I cost/100 g cooked (pasted values, no recipe)
  Κυρίως Πιάτα / Σαλάτες / Πρωινά
                     one row per variant: A dish # · E title (first row of the dish only)
                     · F packaging list · name/weight pairs from G · «Αναλυτικά» (Πρωινά: «Ποικιλία»)
                     · «Prod Cost» · «Pack cost» (located by header text)
"""
import json
import math
import sys

import openpyxl


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and (math.isnan(v) or math.isinf(v))) else float(v)
    s = str(v).strip().replace(',', '.')
    try:
        f = float(s)
        return None if (math.isnan(f) or math.isinf(f)) else f
    except ValueError:
        return None


def txt(v):
    if v is None:
        return None
    s = str(v).replace('\n', ' ').strip()
    return s or None


def raw_flag(v):
    """Return a note for Excel error-ish cells (#REF!, #DIV/0!, Infinity)."""
    if isinstance(v, str) and v.strip().startswith('#'):
        return v.strip()
    if isinstance(v, float) and (math.isinf(v) or math.isnan(v)):
        return 'Infinity/NaN'
    return None


def main(path, out):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    data = {'src_costs': [], 'src_packaging': [], 'index': [], 'groups': [], 'dish_rows': []}

    for i, r in enumerate(wb['Src Costs'].iter_rows(values_only=True)):
        if i == 0 or not txt(r[1]):
            continue
        data['src_costs'].append({'row': i + 1, 'category': txt(r[0]), 'name': txt(r[1]),
                                  'price_kg': num(r[2]), 'err': raw_flag(r[2])})

    for i, r in enumerate(wb['Src Packaging'].iter_rows(values_only=True)):
        if i == 0 or not txt(r[0]):
            continue
        data['src_packaging'].append({'row': i + 1, 'name': txt(r[0]), 'pcs': num(r[1]),
                                      'price_pack': num(r[2]), 'price_pc': num(r[3]),
                                      'note': txt(r[5]) if len(r) > 5 else None,
                                      'err': raw_flag(r[2]) or raw_flag(r[3])})

    for i, r in enumerate(wb['New Index Υλικών'].iter_rows(values_only=True)):
        if i == 0 or not txt(r[1]):
            continue
        data['index'].append({'row': i + 1, 'raw': txt(r[0]), 'name': txt(r[1]),
                              'cost100_raw': num(r[6]), 'cooked_g': num(r[7]), 'cost100_cooked': num(r[8]),
                              'err': raw_flag(r[6]) or raw_flag(r[7]) or raw_flag(r[8])})

    for i, r in enumerate(wb['Group Υλικών'].iter_rows(values_only=True)):
        if i == 0 or not txt(r[1]):
            continue
        data['groups'].append({'row': i + 1, 'name': txt(r[1]), 'cost100_raw': num(r[6]),
                               'cooked_g': num(r[7]), 'cost100_cooked': num(r[8]),
                               'err': raw_flag(r[6]) or raw_flag(r[8])})

    for sheet in ('Κυρίως Πιάτα', 'Σαλάτες', 'Πρωινά'):
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = [txt(c) or '' for c in rows[1]]
        c_an = head.index('Αναλυτικά') if 'Αναλυτικά' in head else head.index('Ποικιλία')
        c_prod = head.index('Prod Cost')
        c_pack = head.index('Pack cost')
        title = None
        last_no = None
        for i, r in enumerate(rows[2:], start=3):
            no = num(r[0])
            if no is None:
                continue
            if no != last_no:
                title = None
            last_no = no
            title = txt(r[4]) or title
            lines = []
            for c in range(6, c_an, 2):
                n, g = txt(r[c]), num(r[c + 1]) if c + 1 < len(r) else None
                if n:
                    lines.append([n, g])
            data['dish_rows'].append({'sheet': sheet, 'row': i, 'dish_no': int(no), 'title': title,
                                      'packaging': txt(r[5]), 'lines': lines, 'analytic': txt(r[c_an]),
                                      'prod_cost': num(r[c_prod]), 'pack_cost': num(r[c_pack])})

    with open(out, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)
    print({k: len(v) for k, v in data.items()})


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
